import io
from datetime import date

import numpy as np
import pymupdf
import pytest
from PIL import Image, ImageDraw
from openpyxl import load_workbook

from lab_extractor.config import get_settings
from lab_extractor.db import Database
from lab_extractor.entities import EntityList
from lab_extractor.exporter import build_workbook
from lab_extractor.image_utils import (
    UnsupportedFile, estimate_skew_angle, load_pages, preprocess,
)
from lab_extractor.pipeline import process_upload
from lab_extractor.schemas import STATUS_APPROVED, STATUS_REVIEW, ExtractedForm
from lab_extractor.validation import normalize_date, normalize_id, split_numbers

S = get_settings()
ENT = EntityList.from_file(S.entities_path)
TODAY = date(2026, 10, 18)


def form(**kw):
    base = dict(is_continuation_page=False, first_name="דנה", last_name="כהן",
                id_number="11223344", central_numbers=["947768"],
                received_date="27-08-2026", sending_entity="שמיר",
                commitment_numbers=["2646"], uncertain_fields=[],
                needs_review=False, review_reason="")
    base.update(kw)
    return ExtractedForm(**base)


class FakeExtractor:
    def __init__(self, forms):
        self.forms, self.i = list(forms), 0

    def extract(self, image_bytes, mime_type="image/jpeg"):
        f = self.forms[self.i]
        self.i += 1
        if isinstance(f, Exception):
            raise f
        return f


def make_pdf(n):
    doc = pymupdf.open()
    for i in range(n):
        page = doc.new_page()
        page.insert_text((72, 72), f"page {i}")
    return doc.tobytes()


# ---------- ת.ז. ותאריכים
def test_id_padding_and_zero_strip():
    assert normalize_id("11223344")[0] == "011223344"
    assert normalize_id("0011223344")[0] == "011223344"
    assert normalize_id("")[0] == ""
    assert normalize_id("AB12345")[1] is not None
    assert normalize_id("1234567890123")[1] is not None


def test_date_normalization():
    assert normalize_date("27.8.26", TODAY) == ("27/08/2026", None)
    assert normalize_date("2026-08-27", TODAY)[0] == "27/08/2026"
    assert normalize_date("31/02/2026", TODAY)[1] is not None
    assert normalize_date("27/08/2030", TODAY)[1] is not None
    assert normalize_date("blah", TODAY)[1] is not None
    assert normalize_date("", TODAY) == ("", None)


def test_split_numbers():
    assert split_numbers(["947768 947769", "947768,12"]) == ["947768", "947769", "12"]


# ---------- גורמים שולחים
@pytest.mark.parametrize("text,expected", [
    ("מכבי שירותי בריאות", "מכבי"),
    ('בי"ח הלל יפה', "הלל יפה"),
    ("המרכז הרפואי שמיר (אסף הרופא)", "אסף הרופא (שמיר)"),
    ('רמב"ם', 'רמב"ם'),
    ("בית חולים ברזילי", "ברזילי"),
])
def test_entity_match(text, expected):
    assert ENT.match(text) == expected


def test_entity_no_match():
    assert ENT.match("בית חולים בדיוני על הירח") is None


# ---------- צינור
def test_pipeline_merge_and_flags():
    pdf = make_pdf(3)
    fake = FakeExtractor([
        form(central_numbers=["947768", "947769"], commitment_numbers=["2646", "2647"]),
        form(is_continuation_page=True, first_name="", last_name="", id_number="",
             central_numbers=[], received_date="", sending_entity="", commitment_numbers=["2648"]),
        form(first_name="", id_number="", uncertain_fields=["id_number"], needs_review=True,
             review_reason="כתב יד לא ברור", sending_entity="לא קיים בשום מקום"),
    ])
    recs = process_upload("f.pdf", pdf, fake, ENT, S, today=TODAY)
    assert len(recs) == 2
    first, second = recs
    assert first.pages == [1, 2] and "1-2" in first.source_file
    assert first.central_number == "947768 947769"
    assert first.commitment_number == "2646 2647 2648"
    assert first.id_number == "011223344" and not first.needs_review
    assert second.needs_review and second.sending_entity == ""
    assert "id_number" in second.uncertain and "first_name" in second.uncertain


def test_pipeline_failure_row_and_bad_file():
    from lab_extractor.extractor import ExtractionError
    recs = process_upload("f.pdf", make_pdf(2), FakeExtractor([form(), ExtractionError("x")]),
                          ENT, S, today=TODAY)
    assert len(recs) == 2 and recs[1].needs_review and "נכשל" in recs[1].reasons[0]
    with pytest.raises(UnsupportedFile):
        process_upload("bad.pdf", b"%PDF-garbage", FakeExtractor([]), ENT, S)
    with pytest.raises(UnsupportedFile):
        process_upload("bad.jpg", b"not an image", FakeExtractor([]), ENT, S)


def test_single_image_has_plain_source_name():
    buf = io.BytesIO()
    Image.new("RGB", (400, 300), "white").save(buf, "PNG")
    recs = process_upload("photo.png", buf.getvalue(), FakeExtractor([form()]), ENT, S, today=TODAY)
    assert recs[0].source_file == "photo.png"


# ---------- מסד נתונים וייצוא
def test_db_and_export(tmp_path):
    db = Database(str(tmp_path / "t.db"))
    recs = process_upload("f.pdf", make_pdf(2),
                          FakeExtractor([form(), form(id_number="", needs_review=True)]),
                          ENT, S, today=TODAY)
    ids = db.insert_many([r.to_row() for r in recs])
    rows = db.fetch_all()
    assert [r["status"] for r in rows][1] == STATUS_REVIEW
    assert db.approve_all_clean() == 1
    db.approve(ids[1])
    assert db.counts()["approved"] == 2
    rows = db.fetch_all()
    data = build_workbook(rows, include_source=True, mark_review=True)
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["C2"].value == "011223344" and ws["D2"].value == "947768"
    assert ws["E2"].value.year == 2026 and ws.sheet_view.rightToLeft
    db.delete([ids[0]])
    assert len(db.fetch_all()) == 1


# ---------- תמונה
def test_deskew_estimates_angle():
    img = Image.new("RGB", (1400, 1000), "white")
    d = ImageDraw.Draw(img)
    for y in range(100, 900, 60):
        d.line([(80, y), (1320, y)], fill="black", width=3)
    rotated = img.rotate(-3, expand=False, fillcolor="white", resample=Image.BICUBIC)
    bgr = np.array(rotated)[:, :, ::-1].copy()
    angle = estimate_skew_angle(bgr)
    assert abs(abs(angle) - 3) < 0.7
    fixed = preprocess(rotated)
    fixed_angle = estimate_skew_angle(np.array(fixed)[:, :, ::-1].copy())
    assert abs(fixed_angle) < abs(angle)


def test_pdf_page_limit():
    with pytest.raises(UnsupportedFile):
        load_pages(make_pdf(3), 1.0, 2)
