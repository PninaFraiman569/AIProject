"""ייצוא לאקסל (.xlsx) בפורמט שהמעבדה עובדת איתו: שורה להתחייבות, עמודה לשדה."""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .schemas import DATA_FIELDS, FIELD_LABELS, LIST_FIELDS, LIST_SEPARATOR, STATUS_APPROVED

_HEADER_FILL = PatternFill("solid", fgColor="14213D")
_REVIEW_FILL = PatternFill("solid", fgColor="FFE8C2")
_THIN = Side(style="thin", color="D5DAE1")


def build_workbook(rows: list[dict], include_source: bool = False,
                   mark_review: bool = False) -> bytes:
    """בונה קובץ אקסל בזיכרון ומחזיר את הבתים שלו.

    ת.ז. ומספרים נשמרים כטקסט (כדי לשמור אפסים מובילים ולהימנע מכתיב מדעי),
    תאריכים נשמרים כתאריכי אקסל אמיתיים. כמה מספרים מרכזיים / מספרי התחייבות
    נכתבים באותו תא, מופרדים בפסיק (״947768, 947769״).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "התחייבויות"
    ws.sheet_view.rightToLeft = True

    headers = [FIELD_LABELS[f] for f in DATA_FIELDS]
    if include_source:
        headers.append("קובץ מקור")
    if mark_review:
        headers.append("סטטוס")
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF", name="Arial")
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, f in enumerate(DATA_FIELDS, start=1):
            value = row.get(f, "") or ""
            if f in LIST_FIELDS:
                value = LIST_SEPARATOR.join(value) if isinstance(value, (list, tuple)) else str(value)
            cell = ws.cell(row=r_idx, column=c_idx)
            if f == "received_date":
                try:
                    cell.value = datetime.strptime(value, "%d/%m/%Y")
                    cell.number_format = "DD/MM/YYYY"
                except ValueError:
                    cell.value = value
                    cell.number_format = "@"
            else:
                cell.value = value
                cell.number_format = "@"
            cell.alignment = Alignment(horizontal="right")
        extra = len(DATA_FIELDS)
        if include_source:
            extra += 1
            ws.cell(row=r_idx, column=extra, value=row.get("source_file", ""))
        if mark_review:
            extra += 1
            ws.cell(row=r_idx, column=extra, value=row.get("status", ""))
            if row.get("status") != STATUS_APPROVED:
                for col in range(1, len(headers) + 1):
                    ws.cell(row=r_idx, column=col).fill = _REVIEW_FILL
        for col in range(1, len(headers) + 1):
            ws.cell(row=r_idx, column=col).border = Border(bottom=_THIN)

    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        longest = max((len(str(c.value or "")) for c in ws[letter]), default=8)
        ws.column_dimensions[letter].width = min(max(longest + 3, 12), 48)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
