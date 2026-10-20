"""בדיקות ללוגיקת הפענוח והנרמול (לא דורשות מסד נתונים או Gemini)."""
import io

import numpy as np
import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from app.core.config import DEFAULT_MODEL, get_settings
from app.core.exporter import build_workbook
from app.core.extractor import ExtractionError
from app.core.image_utils import UnsupportedFile, estimate_skew_angle, load_pages, preprocess
from app.core.pipeline import process_upload
from app.core.validation import normalize_date, normalize_id, normalize_form, split_numbers
from conftest import ENT, S, TODAY, FakeExtractor, form, make_pdf


def test_default_model_is_gemini_38_flash():
    assert DEFAULT_MODEL == "gemini-3.8-flash"


# ---------- ת.ז. ותאריכים
def test_id_padding_and_zero_strip():
    assert normalize_id("11223344")[0] == "011223344"
    assert normalize_id("0011223344")[0] == "011223344"
    assert normalize_id("")[0] == ""
    assert normalize_id("AB12345")[1] is not None
    assert normalize_id("1234567890123")[1] is not None


def test_date_normalization():
    assert normalize_date("27.8.26", TODAY) == ("27/08/2026", None)
    assert normalize_date("2026-08-27", TODAY)[0] == "27/08/2026"
    assert normalize_date("31/02/2026", TODAY)[1] is not None
    assert normalize_date("27/08/2030", TODAY)[1] is not None
    assert normalize_date("blah", TODAY)[1] is not None
    assert normalize_date("", TODAY) == ("", None)


# ---------- כמה מספרים (הבאג מהפרויקט הקודם)
@pytest.mark.parametrize("items,expected", [
    (["947768", "947769"], ["947768", "947769"]),          # רשימה תקינה
    (["947768 947769"], ["947768", "947769"]),             # שני מספרים ברווח
    (["947768, 947769"], ["947768", "947769"]),            # פסיק
    (["947768-947769"], ["947768", "947769"]),             # מקף - היה באג
    (["947768/947769;947770"], ["947768", "947769", "947770"]),
    (["947768\n947769"], ["947768", "947769"]),            # שורה חדשה
    (["947768", "947768"], ["947768"]),                    # כפילות
])
def test_split_numbers_multiple(items, expected):
    assert split_numbers(items)[0] == expected


def test_split_numbers_broken_single_number_is_merged_and_noted():
    numbers, notes = split_numbers(["94776 8"])   # היה מתפצל לשני מספרים שגויים
    assert numbers == ["947768"] and notes


def test_multiple_numbers_kept_in_one_record():
    rec = normalize_form(form(central_numbers=["947768", "947769"],
                              commitment_numbers=["2646-2647"]), 1, ENT, S, TODAY)
    assert rec.central_numbers == ["947768", "947769"]
    assert rec.commitment_numbers == ["2646", "2647"]
    assert not rec.needs_review


# ---------- גורמים שולחים
@pytest.mark.parametrize("text,expected", [
    ("מכבי שירותי בריאות", "מכבי"),
    ('בי"ח הלל יפה', "הלל יפה"),
    ("המרכז הרפואי שמיר (אסף הרופא)", "אסף הרופא (שמיר)"),
    ('רמב"ם', 'רמב"ם'),
    ("בית חולים ברזילי", "ברזילי"),
])
def test_entity_match(text, expected):
    assert ENT.match(text) == expected


def test_entity_no_match():
    assert ENT.match("בית חולים בדיוני על הירח") is None


def test_prompt_explains_service_provider_is_not_sender():
    from app.core.prompts import build_system_prompt
    prompt = build_system_prompt(ENT.names)
    assert "נותן השירות" in prompt and "לעולם אל" in prompt


# ---------- צינור
def test_pipeline_merge_and_flags():
    fake = FakeExtractor([
        form(central_numbers=["947768", "947769"], commitment_numbers=["2646", "2647"]),
        form(is_continuation_page=True, first_name="", last_name="", id_number="",
             central_numbers=[], received_date="", sending_entity="", commitment_numbers=["2648"]),
        form(first_name="", id_number="", uncertain_fields=["id_number"], needs_review=True,
             review_reason="כתב יד לא ברור", sending_entity="לא קיים בשום מקום"),
    ])
    recs = process_upload("f.pdf", make_pdf(3), fake, ENT, S, today=TODAY)
    # הטופס הראשון (עמודים 1-2) עם 3 מספרי התחייבות -> 3 רשומות זהות חוץ מהמספר
    assert len(recs) == 4
    *firsts, second = recs
    assert [r.commitment_numbers for r in firsts] == [["2646"], ["2647"], ["2648"]]
    for first in firsts:
        assert first.pages == [1, 2] and "1-2" in first.source_file
        assert first.central_numbers == ["947768", "947769"]
        assert first.id_number == "011223344" and not first.needs_review
    assert second.needs_review and second.sending_entity == ""
    assert "id_number" in second.uncertain and "first_name" in second.uncertain


def test_pipeline_failure_row_and_bad_file():
    recs = process_upload("f.pdf", make_pdf(2), FakeExtractor([form(), ExtractionError("x")]),
                          ENT, S, today=TODAY)
    assert len(recs) == 2 and recs[1].needs_review and "נכשל" in recs[1].reasons[0]
    with pytest.raises(UnsupportedFile):
        process_upload("bad.pdf", b"%PDF-garbage", FakeExtractor([]), ENT, S)
    with pytest.raises(UnsupportedFile):
        process_upload("bad.jpg", b"not an image", FakeExtractor([]), ENT, S)


def test_parallel_pipeline_keeps_page_order():
    """עם כמה תהליכונים, כל עמוד עדיין מקבל את התוצאה שלו."""
    from dataclasses import replace

    class ByPage:
        def extract(self, image_bytes, mime_type="image/jpeg"):
            return form()

    recs = process_upload("f.pdf", make_pdf(4), ByPage(), ENT,
                          replace(S, max_workers=4), today=TODAY)
    assert [r.pages for r in recs] == [[1], [2], [3], [4]]


def test_single_image_has_plain_source_name():
    buf = io.BytesIO()
    Image.new("RGB", (400, 300), "white").save(buf, "PNG")
    recs = process_upload("photo.png", buf.getvalue(), FakeExtractor([form()]), ENT, S, today=TODAY)
    assert recs[0].source_file == "photo.png"


# ---------- ייצוא
def test_export_joins_multiple_numbers_with_comma():
    rec = normalize_form(form(central_numbers=["947768", "947769"],
                              commitment_numbers=["2646", "2647"]), 1, ENT, S, TODAY)
    row = {**rec.to_row(), "status": "אושר"}
    ws = load_workbook(io.BytesIO(build_workbook([row], include_source=True))).active
    assert ws["C2"].value == "011223344"
    assert ws["D2"].value == "947768, 947769"
    assert ws["G2"].value == "2646, 2647"
    assert ws["E2"].value.year == 2026 and ws.sheet_view.rightToLeft


# ---------- תמונה
def test_deskew_estimates_angle():
    img = Image.new("RGB", (1400, 1000), "white")
    d = ImageDraw.Draw(img)
    for y in range(100, 900, 60):
        d.line([(80, y), (1320, y)], fill="black", width=3)
    rotated = img.rotate(-3, expand=False, fillcolor="white", resample=Image.BICUBIC)
    bgr = np.array(rotated)[:, :, ::-1].copy()
    angle = estimate_skew_angle(bgr)
    assert abs(abs(angle) - 3) < 0.7
    fixed = preprocess(rotated)
    assert abs(estimate_skew_angle(np.array(fixed)[:, :, ::-1].copy())) < abs(angle)


def test_pdf_page_limit():
    with pytest.raises(UnsupportedFile):
        load_pages(make_pdf(3), 1.0, 2)
