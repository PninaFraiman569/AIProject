"""בדיקות לשרת ולמסד (רצות רק כשמוגדר TEST_DATABASE_URL)."""
import io
import time
from dataclasses import replace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import Services, create_app
from app.repository import Repository
from conftest import S, FakeExtractor, form, make_pdf

PASSWORD = "Test-Pass-123"


@pytest.fixture
def client(database_url):
    settings = replace(S, database_url=database_url, app_username="lab",
                       app_password=PASSWORD, gemini_api_key="fake", max_login_attempts=3)
    repo = Repository(database_url)
    repo.init_schema()
    with repo._pool.connection() as conn:   # מסד נקי לכל בדיקה
        conn.execute("TRUNCATE commitments, sessions, users RESTART IDENTITY CASCADE")
    fake = FakeExtractor([
        form(central_numbers=["947768", "947769"], commitment_numbers=["2646", "2647"]),
        form(id_number="", uncertain_fields=["id_number", "received_date"], needs_review=True,
             review_reason="כתב יד לא ברור"),
    ])
    services = Services(settings, repo, extractor_factory=lambda: fake)
    with TestClient(create_app(services)) as c:
        yield c
    repo.close()


def login(client, password=PASSWORD):
    return client.post("/api/auth/login", json={"username": "lab", "password": password})


def auth(client):
    return {"Authorization": f"Bearer {login(client).json()['token']}"}


def wait_job(client, h, job_id):
    for _ in range(100):
        job = client.get(f"/api/jobs/{job_id}", headers=h).json()
        if job["status"] != "running":
            return job
        time.sleep(0.1)
    raise AssertionError("job did not finish")


def test_everything_requires_login(client):
    assert client.get("/api/health").status_code == 200
    for path in ("/api/records", "/api/stats", "/api/export", "/api/config"):
        assert client.get(path).status_code == 401
    assert client.get("/api/records", headers={"Authorization": "Bearer nope"}).status_code == 401


def test_login_logout_and_lockout(client):
    assert login(client, "wrong").status_code == 401
    h = auth(client)
    assert client.get("/api/auth/me", headers=h).json()["username"] == "lab"
    client.post("/api/auth/logout", headers=h)
    assert client.get("/api/auth/me", headers=h).status_code == 401
    for _ in range(3):
        login(client, "wrong")
    r = login(client)  # גם סיסמה נכונה חסומה זמנית
    assert r.status_code == 429


def test_password_is_not_stored_in_plain_text(client):
    login(client)
    repo = client.app.state.svc.repo
    stored = repo.get_user("lab")["password_hash"]
    assert PASSWORD not in stored and stored.startswith("scrypt$")


def test_full_flow_upload_review_edit_export(client):
    h = auth(client)
    files = [("files", ("forms.pdf", make_pdf(2), "application/pdf"))]
    job = client.post("/api/jobs", headers=h, files=files).json()
    job = wait_job(client, h, job["id"])
    assert job["status"] == "done" and job["created"] == 3 and job["progress"] == 1.0

    rows = client.get("/api/records", headers=h).json()
    clean, clean2, flagged = rows
    # שני מספרי התחייבות = שתי רשומות; כמה מספרים מרכזיים נשארים כרשימה באותה רשומה
    assert clean["central_numbers"] == clean2["central_numbers"] == ["947768", "947769"]
    assert clean["commitment_numbers"] == ["2646"] and clean2["commitment_numbers"] == ["2647"]
    assert flagged["needs_review"] and set(flagged["uncertain_fields"]) >= {"id_number", "received_date"}
    assert client.get("/api/stats", headers=h).json()["review"] == 1

    # תיקון תא אחד: הכתום יורד רק ממנו, והשורה עדיין לא מאושרת
    r = client.patch(f"/api/records/{flagged['id']}", headers=h,
                     json={"fields": {"id_number": "12345678"}}).json()
    assert r["id_number"] == "012345678" and r["status"] != "אושר"
    assert "id_number" not in r["uncertain_fields"] and "received_date" in r["uncertain_fields"]

    # ערך לא תקין נדחה עם הודעה
    bad = client.patch(f"/api/records/{flagged['id']}", headers=h,
                       json={"fields": {"received_date": "לא תאריך"}})
    assert bad.status_code == 422 and "תאריך" in bad.json()["detail"]

    # תיקון כל השדות הכתומים -> השורה מאושרת
    for f in r["uncertain_fields"]:
        value = {"received_date": "1.9.26", "first_name": "רחל"}.get(f, "x")
        r = client.patch(f"/api/records/{flagged['id']}", headers=h, json={"fields": {f: value}}).json()
    assert r["status"] == "אושר" and r["uncertain_fields"] == []

    # עריכת מספרים מטקסט עם פסיק
    r = client.patch(f"/api/records/{clean['id']}", headers=h,
                     json={"fields": {"central_numbers": "947768, 947770"}}).json()
    assert r["central_numbers"] == ["947768", "947770"]

    assert client.post("/api/records/approve-clean", headers=h).json()["approved"] == 1
    xlsx = client.get("/api/export?scope=approved&include_source=true", headers=h)
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws.max_row == 3
    assert ws["D2"].value == "947768, 947770"

    assert client.post("/api/records/delete", headers=h, json={"ids": [clean["id"]]}).json()["deleted"] == 1
    assert client.post("/api/records/delete-all", headers=h).json()["deleted"] == 1
    assert client.get("/api/records", headers=h).json() == []


def test_bad_file_reported_not_crashing(client):
    h = auth(client)
    files = [("files", ("bad.pdf", b"%PDF-garbage", "application/pdf"))]
    job = wait_job(client, h, client.post("/api/jobs", headers=h, files=files).json()["id"])
    assert job["status"] == "done" and job["created"] == 0 and "bad.pdf" in job["failed"][0]
